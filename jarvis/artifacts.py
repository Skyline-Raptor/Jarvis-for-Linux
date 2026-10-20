"""Optimistic, recoverable file transactions. No macros are executed."""
import hashlib
import json
import os
import tempfile
import threading
import uuid
from pathlib import Path
from .config import atomic_json


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda:f.read(1024*1024), b''):
            h.update(chunk)
    return h.hexdigest()


class Artifacts:
    def __init__(self, broker):
        self.broker = broker
        self.vault = broker.config.home / 'Versions'
        self.vault.mkdir(exist_ok=True)
        self.lock = threading.RLock()

    def version(self, path):
        p = self.broker.path(path)
        return sha(p) if p.exists() else 'NEW'

    def write(self, path, data, expected):
        p = self.broker.path(path)
        if len(data) > 128*1024*1024:
            raise ValueError('Output exceeds 128 MiB transaction limit; split the output')
        with self.lock:
            if self.version(p) != expected:
                raise RuntimeError('File changed since it was read; read it again before editing')
            p.parent.mkdir(parents=True, exist_ok=True)
            receipt = uuid.uuid4().hex
            if p.exists():
                import shutil
                shutil.copyfile(p, self.vault / (receipt + '.bak'))
            atomic_json(self.vault / (receipt + '.json'), dict(path=str(p), before=expected, after=hashlib.sha256(data).hexdigest()))
            fd, name = tempfile.mkstemp(dir=p.parent, prefix='.jarvis-', suffix=p.suffix)
            try:
                with os.fdopen(fd, 'wb') as f:
                    f.write(data)
                    f.flush()
                    os.fsync(f.fileno())
                if self.version(p) != expected:
                    raise RuntimeError('External edit detected while preparing output')
                os.replace(name, p)
            finally:
                if os.path.exists(name):
                    os.unlink(name)
            if sha(p) != hashlib.sha256(data).hexdigest():
                raise IOError('Output readback mismatch')
            self.broker.store.event('artifact', dict(path=str(p), receipt=receipt, sha256=sha(p)))
            return dict(path=str(p), sha256=sha(p), receipt=receipt, bytes=len(data))

    def restore(self, receipt):
        if not receipt.isalnum() or len(receipt) != 32:
            raise ValueError('Invalid receipt')
        meta = json.loads((self.vault / (receipt + '.json')).read_text(encoding='utf-8'))
        if meta['before'] == 'NEW':
            raise ValueError('This receipt created a new file; no previous content exists')
        return self.write(meta['path'], (self.vault / (receipt + '.bak')).read_bytes(), meta['after'])


def zip_guard(path):
    import zipfile
    with zipfile.ZipFile(path) as z:
        files = z.infolist()
        if len(files) > 10000 or sum(i.file_size for i in files) > 128*1024*1024:
            raise ValueError('Document expands beyond parser limits')
        if any(i.file_size > 32*1024*1024 or i.file_size / max(1,i.compress_size) > 500 for i in files):
            raise ValueError('Suspicious document compression ratio or entry size')


def extract(path):
    p = Path(path)
    if p.stat().st_size > 32*1024*1024:
        raise ValueError('Reader limit is 32 MiB; split large files first')
    ext = p.suffix.lower()
    if ext in ('.docx', '.xlsx', '.pptx'):
        zip_guard(p)
    if ext == '.pdf':
        from pypdf import PdfReader
        r = PdfReader(p)
        if r.is_encrypted:
            raise ValueError('Encrypted PDF: provide an unlocked copy')
        return '\n'.join((page.extract_text() or '') for page in r.pages[:200])[:500000]
    if ext == '.docx':
        from docx import Document
        d = Document(p)
        return '\n'.join([x.text for x in d.paragraphs] + ['\t'.join(c.text for c in row.cells) for t in d.tables for row in t.rows])[:500000]
    if ext == '.xlsx':
        from openpyxl import load_workbook
        w = load_workbook(p, read_only=True, data_only=False)
        try:
            out = []
            for s in w:
                out.append('SHEET: ' + s.title)
                for row in s.iter_rows(max_row=min(s.max_row or 1,10000), max_col=min(s.max_column or 1,100), values_only=True):
                    out.append('\t'.join('' if v is None else str(v) for v in row))
                    if sum(map(len,out)) > 500000:
                        return '\n'.join(out)[:500000]
            return '\n'.join(out)
        finally:
            w.close()
    if ext == '.pptx':
        from pptx import Presentation
        return '\n'.join(shape.text for slide in Presentation(p).slides for shape in slide.shapes if shape.has_text_frame)[:500000]
    if ext in ('.png','.jpg','.jpeg','.webp','.bmp'):
        raise ValueError('Use inspect_image for images')
    data = p.read_bytes()
    if b'\0' in data[:8192]:
        raise ValueError('Unsupported binary format')
    return data.decode('utf-8-sig', errors='strict')[:500000]


def produce(ext, content):
    """Create structured formats from JSON; rich arbitrary files via approved code."""
    from io import BytesIO
    out = BytesIO()
    if ext == '.docx':
        from docx import Document
        d = Document()
        for block in content.get('blocks', []):
            if block.get('heading'):
                d.add_heading(block['text'], min(3,int(block['heading'])))
            else:
                d.add_paragraph(block['text'])
        d.save(out)
    elif ext == '.xlsx':
        from openpyxl import Workbook
        w = Workbook()
        w.remove(w.active)
        for title, rows in content['sheets'].items():
            s = w.create_sheet(title)
            for row in rows:
                s.append(row)
            s.freeze_panes = 'A2'
            s.auto_filter.ref = s.dimensions
        w.save(out)
    elif ext == '.pptx':
        from pptx import Presentation
        r = Presentation()
        for slide in content['slides']:
            s = r.slides.add_slide(r.slide_layouts[1])
            s.shapes.title.text = slide['title']
            s.placeholders[1].text = slide['text']
        r.save(out)
    elif ext == '.pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from xml.sax.saxutils import escape
        styles = getSampleStyleSheet()
        story=[]
        for line in content['text'].splitlines():
            story += [Paragraph(escape(line),styles['Normal']),Spacer(1,8)]
        SimpleDocTemplate(out).build(story)
    else:
        raise ValueError('Structured producer supports docx/xlsx/pptx/pdf')
    return out.getvalue()
